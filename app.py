import streamlit as st
# ------------------------------------------------------------------
# ⚡️ PART 0: SYSTEM CONFIG
# ------------------------------------------------------------------
st.set_page_config(page_title="KKU Scheduler", layout="wide")

import pandas as pd
import io
import math
from ortools.sat.python import cp_model
from openpyxl import Workbook
from openpyxl.styles import Alignment, PatternFill, Border, Side, Font

# ==========================================
# ⚙️ PART 1: CONSTANTS & CONFIG
# ==========================================
class Config:
    DAYS = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri']
    TIME_SLOTS = range(8, 20)
    LUNCH_BREAK_HOUR = 12
    
    MAX_CAPACITY_LAB = 45     
    MAX_CAPACITY_LEC = 100    
    
    # ถ้าไม่มีข้อมูลนักเรียน ให้ใช้ค่า Default นี้
    DEFAULT_STUDENTS = 40 

    COLOR_MAP = {
        'SC': 'FFF59D', 'CP': 'B3E5FC', 'LI': 'C8E6C9', 
        'GE': 'FFE0B2', 'EN': 'E1BEE7', 'DEFAULT': 'F5F5F5'
    }

# ==========================================
# 📦 PART 2: DATA STRUCTURES
# ==========================================
class Course:
    def __init__(self, data):
        self.code = str(data.get('course_code', 'N/A')).strip()
        self.name = str(data.get('course_name', 'Unknown')).strip()
        self.major = data.get('For_Major', 'Gen')
        self.year = int(data.get('year', 1))
        self.program = data.get('program', 'Regular') # Regular / Special
        
        self.type = data.get('type', 'Lec')
        self.duration = int(data.get('duration', 3))
        self.instructor = str(data.get('instructor', 'TBA'))
        self.students = int(data.get('student_count', 40))
        self.section_idx = int(data.get('section_idx', 1))
        
        # Unique ID: แยก Major, Year, Program, Code, Sec
        self.uid = f"{self.major}_{self.year}_{self.program}_{self.code}_{self.type}_S{self.section_idx}_{id(self)}"

    def __repr__(self):
        return f"{self.code} ({self.program}) Sec.{self.section_idx}"

# ==========================================
# 🧠 PART 3: SOLVER ENGINE
# ==========================================
class UniversityScheduler:
    def __init__(self, courses, rooms, fixed_data, options):
        self.courses = courses
        self.rooms = sorted(rooms, key=lambda x: x['capacity'])
        self.fixed_data = fixed_data
        self.options = options 
        self.assignments = {}
        self.failed_courses = []
        
    def solve(self):
        model = cp_model.CpModel()
        shifts = {} 
        scheduled_vars = [] 
        
        active_days = Config.DAYS + (['Sat'] if self.options['allow_saturday'] else [])
        allow_lunch = self.options['allow_lunch']
        capacity_flex = 0.9 if self.options['allow_squeeze'] else 1.0 

        st.write(f"⚙️ กำลังประมวลผล {len(self.courses)} รายการวิชา (รวมทุกภาค/ทุก Sec)...")

        # --- A. Variables ---
        for c in self.courses:
            valid_rooms = []
            
            for r in self.rooms:
                # Logic ความจุห้อง (ยืดหยุ่นได้ถ้าติ๊ก allow_squeeze)
                if r['capacity'] >= (c.students * capacity_flex):
                    if c.type.lower() == 'lab':
                        if 'lab' in r['type'].lower(): valid_rooms.append(r)
                    else:
                        valid_rooms.append(r)

            if not valid_rooms:
                self.failed_courses.append(c)
                continue
            
            # Optimization: เลือกห้องที่ดีที่สุด 10 ห้องแรก
            valid_rooms = valid_rooms[:10]

            c_moves = []
            for d in active_days:
                for h in Config.TIME_SLOTS:
                    if h + c.duration > 20: continue
                    
                    if not allow_lunch:
                        if Config.LUNCH_BREAK_HOUR in range(h, h + c.duration): continue
                    
                    # Check Fixed
                    for r in valid_rooms:
                        is_busy = False
                        for t in range(h, h + c.duration):
                            if (d, t, r['name']) in self.fixed_data:
                                is_busy = True; break
                        if is_busy: continue

                        var = model.NewBoolVar(f'shift_{c.uid}_{d}_{h}_{r["name"]}')
                        shifts[(c.uid, d, h, r["name"])] = var
                        c_moves.append(var)
            
            if c_moves:
                is_scheduled = model.NewBoolVar(f'sched_{c.uid}')
                model.Add(sum(c_moves) == is_scheduled)
                scheduled_vars.append(is_scheduled)
            else:
                 self.failed_courses.append(c)

        # --- B. Constraints ---
        time_map = {}
        # 1. Fixed Data
        for (d, t, r_name), info in self.fixed_data.items():
            key = (d, t)
            if key not in time_map: time_map[key] = []
            time_map[key].append({'type': 'fixed', 'room': r_name})

        # 2. Variables
        for (uid, d, h, r_name), var in shifts.items():
            c = next(x for x in self.courses if x.uid == uid)
            for i in range(c.duration):
                key = (d, h + i)
                if key not in time_map: time_map[key] = []
                # Group Key: Major + Year + Program (ภาคปกไม่ชนปก, พิเศษไม่ชนพิเศษ)
                grp_key = f"{c.major}_{c.year}_{c.program}"
                time_map[key].append({
                    'type': 'var', 'var': var, 'room': r_name, 
                    'grp': grp_key, 'instr': c.instructor
                })

        # Check Conflicts
        for slot, items in time_map.items():
            # Room Conflict
            room_usage = {}
            for item in items:
                r = item['room']
                if r not in room_usage: room_usage[r] = []
                if item['type'] == 'var': room_usage[r].append(item['var'])
                else: room_usage[r].append(1)

            for r, vars in room_usage.items():
                if any(isinstance(v, int) for v in vars):
                    for v in vars: 
                        if not isinstance(v, int): model.Add(v == 0)
                elif len(vars) > 1:
                    model.Add(sum(vars) <= 1)

            # Instructor Conflict (อาจารย์คนเดียว สอนพร้อมกันไม่ได้ แม้คนละภาค)
            instr_usage = {}
            grp_usage = {}
            for item in items:
                if item['type'] == 'var':
                    ins = item.get('instr')
                    if ins and ins != 'TBA':
                        if ins not in instr_usage: instr_usage[ins] = []
                        instr_usage[ins].append(item['var'])
                    
                    grp = item.get('grp')
                    if grp:
                        if grp not in grp_usage: grp_usage[grp] = []
                        grp_usage[grp].append(item['var'])
            
            for _, vars in instr_usage.items():
                if len(vars) > 1: model.Add(sum(vars) <= 1)
            # Student Group Conflict
            for _, vars in grp_usage.items():
                if len(vars) > 1: model.Add(sum(vars) <= 1)

        # --- C. Solve ---
        model.Maximize(sum(scheduled_vars))
        solver = cp_model.CpSolver()
        solver.parameters.num_search_workers = 1 
        solver.parameters.max_time_in_seconds = 600.0
        solver.parameters.log_search_progress = True
        
        status = solver.Solve(model)

        if status in [cp_model.OPTIMAL, cp_model.FEASIBLE]:
            for (uid, d, h, r_name), var in shifts.items():
                if solver.Value(var) == 1:
                    c = next(course for course in self.courses if course.uid == uid)
                    self.assignments[c] = (d, h, r_name)
            
            assigned_uids = {c.uid for c in self.assignments.keys()}
            for c in self.courses:
                if c.uid not in assigned_uids and c not in self.failed_courses:
                    self.failed_courses.append(c)
            
            self.failed_courses = list(set(self.failed_courses))
            return True
        else:
            return False

# ==========================================
# 📊 PART 4: EXCEL GENERATOR
# ==========================================
def generate_excel(assignments, active_days):
    output = io.BytesIO()
    wb = Workbook()
    wb.remove(wb.active)
    
    sheet_data = {}
    for c, (d, t, r) in assignments.items():
        # แยก Sheet ตาม Major + Year + Program
        # ex: CS_Y1_Regular, CS_Y1_Special
        key = f"{c.major}_Y{c.year}_{c.program}"
        if key not in sheet_data: sheet_data[key] = []
        sheet_data[key].append((c, d, t, r))
        
    thin = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    
    for sheet_name in sorted(sheet_data.keys()):
        safe_name = sheet_name.replace('/', '_')[:30]
        ws = wb.create_sheet(title=safe_name)
        
        ws.merge_cells('A1:N1')
        ws['A1'] = f"ตารางเรียน {sheet_name}"
        ws['A1'].font = Font(size=14, bold=True)
        ws['A1'].alignment = align
        
        ws['A2'] = "Day/Time"
        for i, h in enumerate(Config.TIME_SLOTS):
            col = chr(66 + i) 
            ws[f'{col}2'] = f"{h}:00"
            ws.column_dimensions[col].width = 16
            
        row_map = {d: i+3 for i, d in enumerate(active_days)}
        for d, r_idx in row_map.items():
            ws[f'A{r_idx}'] = d
            ws[f'A{r_idx}'].font = Font(bold=True)
            for col_idx in range(2, 15):
                cell = ws.cell(row=r_idx, column=col_idx)
                cell.border = thin

        for c, d, t, r in sheet_data[sheet_name]:
            if d not in row_map: continue
            r_idx = row_map[d]
            for i in range(c.duration):
                h_cur = t + i
                if h_cur in Config.TIME_SLOTS:
                    c_idx = Config.TIME_SLOTS.index(h_cur) + 2
                    cell = ws.cell(row=r_idx, column=c_idx)
                    sec_info = f" (Sec {c.section_idx})" if c.section_idx > 1 else ""
                    cell.value = f"{c.code}{sec_info}\n{c.name}\n{r} ({c.instructor})"
                    cell.alignment = align
                    cell.border = thin
                    color = Config.COLOR_MAP.get(c.code[:2].upper(), 'F5F5F5')
                    cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')

    wb.save(output)
    output.seek(0)
    return output

# ==========================================
# 🖥️ PART 5: USER INTERFACE
# ==========================================
st.title("🎓 KKU Scheduler (Fixed Missing Courses)")

with st.sidebar:
    st.header("1. Upload Files")
    f_rooms = st.file_uploader("Rooms (csv)", type=['csv'])
    f_students = st.file_uploader("Students (csv)", type=['csv'])
    f_subjects = st.file_uploader("Subjects (csv)", type=['csv'], accept_multiple_files=True)
    f_fixed = st.file_uploader("Fixed Data (Optional)", type=['csv'], accept_multiple_files=True)
    
    st.markdown("---")
    st.header("2. Survival Options")
    opt_saturday = st.checkbox("📅 เปิดสอนวันเสาร์", value=False)
    opt_lunch = st.checkbox("🍱 เรียนพักเที่ยงได้", value=False)
    opt_squeeze = st.checkbox("🪑 นั่งเบียดได้ 10%", value=True)

if st.button("🚀 Start Scheduling", type="primary"):
    if f_rooms and f_subjects: # Students file is now optional-ish (but recommended)
        with st.spinner("⏳ Processing..."):
            
            # 1. Rooms
            df_rooms = pd.read_csv(f_rooms)
            rooms_list = []
            for _, r in df_rooms.iterrows():
                rooms_list.append({
                    'name': str(r['room_name']),
                    'capacity': int(r['capacity']),
                    'type': str(r['type']).lower()
                })
            
            # 2. Students (สร้าง Map เพื่อดูว่าสาขานี้มีภาคอะไรบ้าง)
            std_map = {} 
            prog_map = {} 
            if f_students:
                df_std = pd.read_csv(f_students)
                for _, r in df_std.iterrows():
                    try:
                        maj, yr, prog = r['major'], int(r['year']), r['program']
                        std_map[(maj, yr, prog)] = int(r['student_count'])
                        
                        if (maj, yr) not in prog_map: prog_map[(maj, yr)] = []
                        if prog not in prog_map[(maj, yr)]: prog_map[(maj, yr)].append(prog)
                    except: pass
            
            # 3. Subjects (หัวใจสำคัญ: บังคับสร้างวิชาแม้ไม่มีข้อมูลนักเรียน)
            all_courses = []
            for f in f_subjects:
                df = pd.read_csv(f)
                maj = f.name.split('_')[1].split('.')[0] 
                
                for _, row in df.iterrows():
                    yr = int(row['year'])
                    
                    # 🔎 เช็ค: ปีนี้มีภาคอะไรบ้าง? (Regular/Special)
                    # ถ้าหาไม่เจอในไฟล์นักเรียน -> ให้ Default เป็น ['Regular'] ไว้ก่อน (กันวิชาหาย)
                    available_progs = prog_map.get((maj, yr), ['Regular'])
                    
                    # ลูปสร้างวิชาให้ครบทุกภาค (Reg/Sp)
                    for prog in available_progs:
                        # ถ้าไม่ทราบจำนวนคน -> ใช้ Default (40)
                        total_students = std_map.get((maj, yr, prog), Config.DEFAULT_STUDENTS)
                        
                        def create_course_variants(c_type, hours):
                            if hours <= 0: return
                            limit = Config.MAX_CAPACITY_LAB if c_type == 'Lab' else Config.MAX_CAPACITY_LEC
                            effective_limit = int(limit * 1.1) if opt_squeeze else limit
                            
                            num_secs = 1
                            if total_students > effective_limit:
                                num_secs = math.ceil(total_students / effective_limit)
                            
                            count_per_sec = math.ceil(total_students / num_secs)
                            
                            for i in range(num_secs):
                                d = row.to_dict()
                                d.update({
                                    'For_Major': maj, 'program': prog, 
                                    'student_count': count_per_sec, 'type': c_type,
                                    'duration': hours, 'section_idx': i + 1
                                })
                                all_courses.append(Course(d))

                        create_course_variants('Lec', row['lecture_hours'])
                        create_course_variants('Lab', row['lab_hours'])

            # 4. Fixed Data
            fixed_data = {}
            if f_fixed:
                for f in f_fixed:
                    dfx = pd.read_csv(f)
                    for _, r in dfx.iterrows():
                        try:
                            s, e = int(float(str(r['start_time']).split(':')[0])), int(float(str(r['end_time']).split(':')[0]))
                            for t in range(s, e): fixed_data[(str(r['day']), t, str(r['room']))] = 'FIXED'
                        except: pass

            # 5. Solve
            user_options = {
                'allow_saturday': opt_saturday,
                'allow_lunch': opt_lunch,
                'allow_squeeze': opt_squeeze
            }
            
            scheduler = UniversityScheduler(all_courses, rooms_list, fixed_data, user_options)
            success = scheduler.solve()
            
            if success:
                st.balloons()
                total_c = len(all_courses)
                success_c = len(scheduler.assignments)
                failed_c = len(scheduler.failed_courses)
                
                if failed_c == 0:
                    st.success(f"🎉 จัดตารางเสร็จสมบูรณ์! ({success_c} รายการ)")
                else:
                    st.warning(f"⚠️ จัดได้ {success_c}/{total_c} รายการ (หาลงไม่ได้ {failed_c} รายการ)")
                
                active_days = Config.DAYS + (['Sat'] if opt_saturday else [])
                excel_file = generate_excel(scheduler.assignments, active_days)
                st.download_button("📥 ดาวน์โหลดตารางเรียน (Excel)", excel_file, "Schedule.xlsx")
                
                if scheduler.failed_courses:
                    with st.expander(f"❌ รายชื่อวิชาที่หลุด ({failed_c})"):
                        for c in scheduler.failed_courses:
                            st.write(f"- {c} (Major: {c.major}, Year: {c.year}) -> ต้องการ {c.students} ที่นั่ง")
            else:
                st.error("❌ Solver Timeout: ลองลดจำนวนวิชา หรือเพิ่มตัวช่วยใน Sidebar")
    else:
        st.info("กรุณาอัปโหลดไฟล์ Rooms และ Subjects เป็นอย่างน้อย")
